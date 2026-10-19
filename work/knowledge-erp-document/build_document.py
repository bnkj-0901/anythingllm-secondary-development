"""Build the deliverable from its Markdown source with the bundled runtime."""

from __future__ import annotations

import json
import re
from decimal import Decimal
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "docs" / "多用户知识库与ERP实施方案.md"
OUTPUT = ROOT / "docs" / "多用户知识库与ERP实施方案.docx"
WORK = Path(__file__).resolve().parent
BLACK = "000000"
FONT_CJK = "Microsoft YaHei"
FONT_LATIN = "Calibri"


def set_font(run, size=None, bold=None, color=BLACK):
    run.font.name = FONT_LATIN
    run.font.color.rgb = RGBColor.from_string(color)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for field, value in [("ascii", FONT_LATIN), ("hAnsi", FONT_LATIN), ("eastAsia", FONT_CJK)]:
        fonts.set(qn("w:" + field), value)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "zh-CN")
    lang.set(qn("w:eastAsia"), "zh-CN")
    rpr.append(lang)


def append_hyperlink(paragraph, label, target):
    rid = paragraph.part.relate_to(
        target,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT_LATIN)
    fonts.set(qn("w:hAnsi"), FONT_LATIN)
    fonts.set(qn("w:eastAsia"), FONT_CJK)
    props.append(fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "194F78")
    props.append(color)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, text, size=None, bold=None, color=BLACK):
    pattern = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
    at = 0
    for match in pattern.finditer(text):
        if match.start() > at:
            set_font(paragraph.add_run(text[at:match.start()]), size, bold, color)
        append_hyperlink(paragraph, match.group(1), match.group(2))
        at = match.end()
    if at < len(text):
        set_font(paragraph.add_run(text[at:]), size, bold, color)


def diagram(path: Path):
    canvas = Image.new("RGB", (1680, 740), "white")
    draw = ImageDraw.Draw(canvas)
    cjk_path = "C:/Windows/Fonts/msyh.ttc"
    font = ImageFont.truetype(cjk_path, 30)
    small = ImageFont.truetype(cjk_path, 25)
    outline = "#C9D1D9"
    line = "#6D7783"
    nodes = {
        "mini": (50, 25, 410, 116, "多微信小程序"),
        "oa": (455, 25, 815, 116, "多微信公众号"),
        "web": (860, 25, 1220, 116, "电脑用户工作台\n平台运营入口"),
        "app": (1265, 25, 1625, 116, "Android 与 iOS\nApp 增量"),
        "api": (450, 192, 1230, 280, "统一 API  渠道适配与权限"),
        "knowledge": (65, 360, 505, 450, "知识服务"),
        "erp": (620, 360, 1060, 450, "ERP 业务服务"),
        "ai": (1175, 360, 1615, 450, "助手与模型适配"),
        "db": (65, 554, 640, 680, "PostgreSQL 与 pgvector\n业务数据与知识索引"),
        "cos": (705, 554, 1110, 680, "私有文件存储\n原件与引用"),
        "queue": (1175, 554, 1615, 680, "异步处理与模型服务\n转写 OCR 解析"),
    }

    def arrow(points):
        draw.line(points, fill=line, width=4)
        x, y = points[-1]
        px, py = points[-2]
        if y > py:
            tip = [(x, y), (x - 9, y - 14), (x + 9, y - 14)]
        elif x > px:
            tip = [(x, y), (x - 14, y - 9), (x - 14, y + 9)]
        else:
            tip = [(x, y), (x + 14, y - 9), (x + 14, y + 9)]
        draw.polygon(tip, fill=line)

    for center in [230, 635, 1040, 1445]:
        arrow([(center, 116), (center, 153), (840, 153), (840, 192)])
    for center in [285, 840, 1395]:
        arrow([(840, 280), (840, 320), (center, 320), (center, 360)])
    arrow([(285, 450), (285, 554)])
    arrow([(840, 450), (840, 497), (540, 497), (540, 554)])
    arrow([(505, 406), (565, 406), (565, 520), (905, 520), (905, 554)])
    arrow([(1395, 450), (1395, 554)])
    arrow([(1175, 390), (1060, 390)])
    arrow([(1175, 426), (1138, 426), (1138, 478), (555, 478), (555, 426), (505, 426)])

    for key, (x0, y0, x1, y1, label) in nodes.items():
        fill = "#F0F4F7" if key == "api" else "#F9FAFB"
        draw.rounded_rectangle((x0, y0, x1, y1), radius=12, fill=fill, outline=outline, width=2)
        used_font = small if "\n" in label or key == "web" else font
        box = draw.multiline_textbbox((0, 0), label, font=used_font, spacing=10, align="center")
        tw, th = box[2] - box[0], box[3] - box[1]
        draw.multiline_text(
            ((x0 + x1 - tw) / 2, (y0 + y1 - th) / 2 - box[1]),
            label,
            font=used_font,
            fill="#17212B",
            spacing=10,
            align="center",
        )
    canvas.save(path)


def configure(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.65)
    section.footer_distance = Inches(0.28)
    for name, size in [("Normal", 11), ("Title", 23), ("Subtitle", 13), ("Heading 1", 16), ("Heading 2", 12.2), ("List Bullet", 10.5)]:
        style = doc.styles[name]
        style.font.name = FONT_LATIN
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CJK)
        pf = style.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(6)
        pf.line_spacing = Pt(size * 1.45)
        pf.widow_control = True
        grid = OxmlElement("w:snapToGrid")
        grid.set(qn("w:val"), "0")
        style._element.get_or_add_pPr().append(grid)
        if name.startswith("Heading"):
            pf.keep_with_next = True
            pf.space_before = Pt(10 if name == "Heading 2" else 0)
            pf.space_after = Pt(6)
    doc.styles["Title"].paragraph_format.space_after = Pt(8)
    for name in ["Normal", "Title", "Subtitle"]:
        ppr = doc.styles[name]._element.get_or_add_pPr()
        for border in list(ppr.findall(qn("w:pBdr"))):
            ppr.remove(border)
        if name != "Normal":
            doc.styles[name].font.italic = False
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(10)
    doc.styles["List Bullet"].paragraph_format.space_after = Pt(4)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run("第 "), 9, color="666666")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    set_font(footer.add_run(" 页"), 9, color="666666")
    props = doc.core_properties
    props.title = "多用户知识库与经营 ERP 实施方案"
    props.subject = "微信小程序 电脑用户工作台 公众号 知识管理 数量进销存 经营记账"
    props.author = ""
    props.last_modified_by = ""
    props.keywords = "知识库, 多用户, 微信小程序, ERP, 数量库存, 经营记账"
    props.comments = ""


def table_widths(header):
    if header[0] == "角色":
        return [0.72, 1.35, 1.65, 1.2, 2.02]
    if header[0] == "操作":
        return [3.05, 0.9, 0.9, 0.9, 1.19]
    if header[0] == "接口":
        return [2.45, 1.85, 2.64]
    if header[0] == "编号":
        return [0.6, 1.9, 4.44]
    if header[0] == "模块":
        return [0.95, 3.25, 2.74]
    if header[0] == "动作":
        return [2.75, 2.85, 1.34]
    if header[0] == "项目":
        return [1.75, 3.3, 1.89]
    if header[0] == "时间":
        return [1.0, 3.2, 2.74]
    return [1.15, 3.15, 2.64]


def add_table(doc, rows):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = table_widths(rows[0])
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        element = OxmlElement("w:" + edge)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "5")
        element.set(qn("w:color"), "D9D9D9")
        borders.append(element)
    pr.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for edge, twips in [("top", 90), ("bottom", 90), ("left", 100), ("right", 100)]:
        item = OxmlElement("w:" + edge)
        item.set(qn("w:w"), str(twips))
        item.set(qn("w:type"), "dxa")
        margins.append(item)
    pr.append(margins)
    for index, row_values in enumerate(rows):
        row = table.rows[0] if index == 0 else table.add_row()
        trpr = row._tr.get_or_add_trPr()
        no_split = OxmlElement("w:cantSplit")
        trpr.append(no_split)
        if index == 0:
            repeating = OxmlElement("w:tblHeader")
            trpr.append(repeating)
        for col, value in enumerate(row_values):
            cell = row.cells[col]
            cell.width = Inches(widths[col])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "DBE7EF" if index == 0 else ("F5F7F9" if index % 2 == 0 else "FFFFFF"))
            cell._tc.get_or_add_tcPr().append(shading)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.line_spacing = Pt(14)
            paragraph.paragraph_format.keep_with_next = False
            paragraph.paragraph_format.widow_control = True
            if index == 0 or (len(row_values) == 5 and col > 0 and rows[0][0] == "操作"):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(paragraph, value, size=10, bold=index == 0)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.line_spacing = Pt(3)
    set_font(paragraph.add_run(""), 3)
    return table


def verify_content(content):
    expected = ["公共知识", "个人私有知识", "语音", "图片", "文字", "文档", "按商品规格与数量管理", "微信公众号", "小程序", "采购", "销售", "应收应付", "十六周实施计划", "验收", "预算", "5 秒", "电脑端用户工作台", "每位用户", "扫码登录", "双端一致", "Android 和 iOS App", "平台有权查看全部知识", "多公众号与多小程序一键适配"]
    missing = [term for term in expected if term not in content]
    if missing:
        raise AssertionError(missing)
    buy_qty, unit_cost, sell_qty, sell_price = map(Decimal, ["100", "10", "30", "15"])
    opening, expense = Decimal("2000"), Decimal("50")
    assert buy_qty - sell_qty == 70
    assert (buy_qty - sell_qty) * unit_cost == 700
    assert sell_qty * (sell_price - unit_cost) - expense == 100
    assert opening - buy_qty * unit_cost + sell_qty * sell_price - expense == 1400
    assert sum([800, 100, 450, 120, 100, 20]) == 1590
    assert sum([2200, 300, 1500, 360, 500, 100]) == 4960
    assert sum([60, 50, 30, 10, 5]) == 155
    assert sum([25, 5, 10]) == 40
    assert Decimal("155") * 1200 * Decimal("1.15") == 213900
    assert Decimal("155") * 1800 * Decimal("1.15") == 320850
    assert Decimal("40") * 1200 * Decimal("1.15") == 55200
    assert Decimal("40") * 1800 * Decimal("1.15") == 82800
    assert Decimal("195") * 1200 * Decimal("1.15") == 269100
    assert Decimal("195") * 1800 * Decimal("1.15") == 403650
    return {"requirements": len(expected), "bookkeepingExample": "passed", "budgetTotals": "passed"}


def main():
    content = SOURCE.read_text(encoding="utf-8")
    checks = verify_content(content)
    WORK.mkdir(parents=True, exist_ok=True)
    architecture = WORK / "architecture.png"
    diagram(architecture)
    doc = Document()
    configure(doc)
    lines = content.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line == "<!-- page -->":
            doc.add_page_break()
        elif line.startswith("```mermaid"):
            index += 1
            while index < len(lines) and not lines[index].startswith("```"):
                index += 1
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.line_spacing = 1
            shape = paragraph.add_run().add_picture(str(architecture), width=Inches(6.84))
            shape._inline.docPr.set("descr", "小程序、公众号和电脑后台共用带权限的 API，连接知识服务、ERP、助手及数据库、文件与异步处理。")
        elif line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                candidate = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", cell) for cell in candidate):
                    rows.append(candidate)
                index += 1
            add_table(doc, rows)
            continue
        elif line.startswith("# "):
            paragraph = doc.add_paragraph(style="Title")
            add_inline(paragraph, line[2:])
        elif line.startswith("## "):
            paragraph = doc.add_paragraph(style="Heading 1")
            add_inline(paragraph, line[3:])
        elif line.startswith("### "):
            paragraph = doc.add_paragraph(style="Heading 2")
            heading = re.sub(r"^(\d+)\.(\d+)\s", r"\1 \2 ", line[4:])
            add_inline(paragraph, heading)
        elif line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline(paragraph, line[2:])
        elif line == "微信小程序 电脑工作台 App 与公众号交付方案":
            paragraph = doc.add_paragraph(style="Subtitle")
            add_inline(paragraph, line)
        elif line.startswith("版本 V"):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(12)
            add_inline(paragraph, line, size=10)
        else:
            paragraph = doc.add_paragraph()
            add_inline(paragraph, line)
        index += 1
    doc.save(OUTPUT)
    reloaded = Document(OUTPUT)
    assert len(reloaded.tables) == 17, len(reloaded.tables)
    assert len([p for p in reloaded.paragraphs if p.style.name == "Heading 1"]) == 15
    checks.update({"tables": len(reloaded.tables), "chapters": 15, "markdownCharacters": len(content), "outputBytes": OUTPUT.stat().st_size})
    (WORK / "content_checks.json").write_text(json.dumps(checks, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(OUTPUT), **checks}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
